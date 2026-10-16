from stage2_setup import *
import pandas as pd,numpy as np,math,ast
from openpyxl import load_workbook
from PIL import Image,ImageDraw
from pypdf import PdfReader
from statsmodels.stats.multitest import multipletests
payload=json.loads((EV/'workbook_payload.json').read_text(encoding='utf-8'));cell_count=0;sheetchecks=[]
for file,sheets in payload.items():
 wb=load_workbook(OUT/file,read_only=True,data_only=False);assert wb.sheetnames==[s['name'] for s in sheets]
 for s in sheets:
  rows=list(wb[s['name']].iter_rows(min_row=5,max_row=5+len(s['rows']),max_col=len(s['columns']),values_only=True));assert list(rows[0])==[c.replace('_',' ') for c in s['columns']]
  for i,row in enumerate(s['rows']):
   for j,c in enumerate(s['columns']):
    a=rows[i+1][j];e=row.get(c);cell_count+=1
    if e is None:assert a is None,(file,s['name'],i,c,a)
    elif isinstance(e,bool):assert a==e
    elif isinstance(e,(int,float)):assert isinstance(a,(int,float)) and math.isclose(a,e,rel_tol=1e-11,abs_tol=1e-9),(file,s['name'],i,c,a,e)
    else:assert (a is None and e=='') or str(a).replace('\r\n','\n')==str(e).replace('\r\n','\n'),(file,s['name'],i,c,a,e)
  sheetchecks.append(dict(file=file,sheet=s['name'],data_rows=len(s['rows']),cells_verified=len(s['rows'])*len(s['columns'])))
 wb.close()
read=lambda n:pd.read_csv(EV/(n+'.csv'),low_memory=False)
c=read('coefficients');m=read('model_metadata');memb=read('sample_membership');groups=read('group_support');b=read('wild_bootstrap');p=read('PPML_results');checks=read('implementation_checks');repl=read('replication_discrepancy_summary');summary=read('priority_summary')
assert len(m)==22 and m.model.nunique()==22;assert len(b)==4 and b.valid_draws.eq(9999).all() and b.seed.eq(20261016).all();assert p.model.nunique()==6 and p.converged.all() and p.positive_rank.eq(p['rank']).all() and p.relative_score.lt(1e-6).all();assert repl.passed.all()
assert (m.n-m['rank']).eq(m.residual_df).all();assert not memb.duplicated(['model','country_code','year']).any()
for label,n,k in [('A primary',914,109),('B primary',1094,135),('C primary',1094,136)]:
 r=m[m.model.eq(label)].iloc[0];assert (r.n,r.clusters)==(n,k)
 a=memb[memb.model.eq(label)];bb=memb[memb.model.eq(label[0]+' matched baseline')];assert set(zip(a.country_code,a.year))==set(zip(bb.country_code,bb.year))
 assert set(groups.loc[groups.model.eq(label),'income_group'])==set(['Pooled','Low income','Lower middle income','Upper middle income','High income'])
assert checks.max_coefficient_discrepancy.lt(1e-8).all() and checks.max_residual_discrepancy.lt(1e-7).all()
assert np.allclose(summary.p_Holm_primary3,multipletests(summary.p_value,method='holm')[1]);assert np.allclose(summary.p_wild_Holm_primary3,multipletests(summary.p_wild_restricted,method='holm')[1]);assert summary.CI95_low.lt(0).all() and summary.CI95_high.gt(0).all()
for _,r in m.iterrows():
 a=memb[memb.model.eq(r.model)];assert len(a)==r.n and a.country_code.nunique()==r.clusters
 assert a.year.between(2014,2023).all()
 assert a.groupby('country_code').size().ge(2).all() and a.groupby('year').size().ge(2).all()
protected=json.loads((EV/'protected_before.json').read_text(encoding='utf-8'));assert all(sha(Path(f))==s for f,s in protected.items())
old=json.loads((ROOT/'reports/irena2025_api_evidence/preexisting_files_manifest.json').read_text(encoding='utf-8'));assert all(Path(r['path']).exists() and sha(Path(r['path']))==r['sha256'] for r in old)
# Previously validated top-level original output hashes and fsQCA outputs.
original_outputs=[]
for folder in [ROOT/'econometric_results',ROOT/'fsqca_preparation_20261016']:
 j=json.loads((folder/'audit_evidence/final_validation.json').read_text(encoding='utf-8'));out=j['outputs']
 if isinstance(out,dict):
  for f,value in out.items():
   expected=value.get('sha256') if isinstance(value,dict) else value
   candidate=Path(f) if Path(f).is_absolute() else folder/f
   if isinstance(expected,str) and len(expected)==64:assert candidate.exists() and sha(candidate)==expected;original_outputs.append(str(candidate))
figures=[]
for f in FIG.glob('*.png'):
 img=Image.open(f);assert img.width>=1000 and img.height>=500;pdf=f.with_suffix('.pdf');assert len(PdfReader(pdf).pages)==1;figures.append(dict(PNG=f.name,width=img.width,height=img.height,PDF=pdf.name))
assert len(figures)==4
previews=[]
for file,sheets in payload.items():
 tag='regression' if file.startswith('second_stage_regression') else 'matched' if file.startswith('matched') else 'robustness'
 import re
 previews.extend(EV/('preview_'+tag+'_'+re.sub('[^a-zA-Z0-9]','_',s['name'])+'.png') for s in sheets)
assert len(previews)==len(sheetchecks) and all(p.exists() for p in previews)
for tag in ['regression','matched','robustness']:
 files=sorted(p for p in previews if p.name.startswith('preview_'+tag+'_'));w,h=700,220;canvas=Image.new('RGB',(1400,math.ceil(len(files)/2)*h),'#e8edf2');dr=ImageDraw.Draw(canvas)
 for i,f in enumerate(files):
  img=Image.open(f).convert('RGB');img.thumbnail((w-12,h-30));x=i%2*w+6;y=i//2*h+25;canvas.paste(img,(x,y));dr.text((x,y-20),f.stem.replace('preview_'+tag+'_',''),fill='#203047')
 canvas.save(EV/('contact_sheet_'+tag+'.png'))
for name in ['second_stage_econometric_report.md','independent_replication_second_stage.md','execution_instructions.md']:assert (OUT/name).read_text(encoding='utf-8')
for src in Path(__file__).parent.glob('stage2_*'):
 if src.suffix not in ['.py','.R']:continue
 dst=SCRIPTS/src.name
 if src.resolve()!=dst.resolve():shutil.copy2(src,dst)
 if src.suffix=='.py':ast.parse(src.read_text(encoding='utf-8'))
validation=dict(status='PASSED',timestamp_utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),OLS_models=22,CR2=True,independent_R_coefficients=30,independent_R_joint_tests=4,independent_R_no_material_discrepancies=True,restricted_wild_bootstrap_tests=4,draws_per_test=9999,PPML_models=6,PPML_inference='CR1 only; nonlinear CR2/bootstrap not implemented',workbooks=3,sheets=len(sheetchecks),cells_verified=cell_count,PNG_figures=4,PDF_figures=4,source_and_protected_files_unchanged=len(protected),earlier_archive_files_unchanged=len(old),prior_outputs_verified=original_outputs,fsQCA_run=False,financing_vintages_spliced=False,deliverables={f:sha(OUT/f) for f in list(payload)+['second_stage_econometric_report.md','independent_replication_second_stage.md','execution_instructions.md']},sheet_checks=sheetchecks,figure_checks=figures)
(EV/'final_validation.json').write_text(json.dumps(validation,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in validation.items() if k not in ['deliverables','sheet_checks','figure_checks','prior_outputs_verified']},indent=2),flush=True)
